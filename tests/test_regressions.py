#!/usr/bin/env python3
"""Regression tests for Word formatting bugs that must remain closed."""

from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
if str(SCRIPTS) not in sys.path:
    sys.path.insert(0, str(SCRIPTS))

import format_docx  # noqa: E402
import validate_docx  # noqa: E402


def load_rules():
    return json.loads((ROOT / "config" / "rules.json").read_text(encoding="utf-8"))


class WordSkillRegressionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.rules = load_rules()
        cls.template_path = ROOT / cls.rules["template_path"]

    def test_merged_cells_are_visited_once_without_skipping_distinct_cells(self):
        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).merge(table.cell(0, 1))

        paragraphs = list(format_docx.iter_table_paragraphs(table))

        # One merged top cell + two distinct bottom cells.
        self.assertEqual(len(paragraphs), 3)
        self.assertEqual(len({p._p for p in paragraphs}), 3)

    def test_style_sync_restores_based_on_after_all_dependencies_exist(self):
        template = Document(self.template_path)
        target = Document()

        format_docx.sync_template_styles(target, template, self.rules)

        for name in self.rules["styles"]["sync_from_template"]:
            source = template.styles[name]
            synced = target.styles[name]
            source_base = source.base_style.name if source.base_style is not None else None
            target_base = synced.base_style.name if synced.base_style is not None else None
            self.assertEqual(
                target_base,
                source_base,
                f"basedOn mismatch after sync for style {name!r}",
            )

    def test_footer_fix_matches_template_semantics_and_removes_extra_first_footer(self):
        template = Document(self.template_path)
        target = Document()

        section = target.sections[-1]
        section.different_first_page_header_footer = True
        section.footer.paragraphs[0].text = "WRONG DEFAULT"
        section.even_page_footer.paragraphs[0].text = "WRONG EVEN"
        section.first_page_footer.paragraphs[0].text = "EXTRA FIRST"

        format_docx.copy_footers(target, template)

        expected_types = format_docx.footer_reference_types(template.sections[-1])
        actual_types = format_docx.footer_reference_types(target.sections[-1])
        self.assertEqual(actual_types, expected_types)
        self.assertEqual(
            target.sections[-1].different_first_page_header_footer,
            template.sections[-1].different_first_page_header_footer,
        )

        with tempfile.TemporaryDirectory() as tmp:
            target_path = Path(tmp) / "target.docx"
            target.save(target_path)

            target_pkg = validate_docx.Package(target_path)
            template_pkg = validate_docx.Package(self.template_path)
            try:
                target_doc = target_pkg.xml("word/document.xml")
                template_doc = template_pkg.xml("word/document.xml")
                self.assertEqual(
                    validate_docx.footer_signature(target_pkg, target_doc),
                    validate_docx.footer_signature(template_pkg, template_doc),
                )
            finally:
                target_pkg.close()
                template_pkg.close()


if __name__ == "__main__":
    unittest.main()
