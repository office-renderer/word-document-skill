#!/usr/bin/env python3
"""Safely normalize Word body/table formatting without reserializing OOXML with ElementTree.

This script deliberately uses python-docx/lxml so existing OOXML namespaces, mc:Ignorable
prefixes, relationships, fields, drawings, bookmarks, and other unsupported elements remain
on their original XML trees. It then performs package/XML/reopen checks before replacing
an in-place document.
"""

from __future__ import annotations

import argparse
import os
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt
from lxml import etree

BODY_STYLE = "正文（默认）"
TABLE_STYLE = "表格正文"
BODY_FIRST_LINE_PT = 32
TABLE_FONT_PT = 14
TWIP_EMU = 635


def set_pagination_off(fmt) -> None:
    """Uncheck all four Word line/page-break paragraph options."""
    fmt.keep_with_next = False
    fmt.keep_together = False
    fmt.page_break_before = False
    fmt.widow_control = False


def set_indentation_zero(paragraph_or_style) -> None:
    """Explicitly block both twip- and character-based indent inheritance."""
    fmt = paragraph_or_style.paragraph_format
    fmt.left_indent = Pt(0)
    fmt.right_indent = Pt(0)
    fmt.first_line_indent = Pt(0)

    if hasattr(paragraph_or_style, "_p"):
        ppr = paragraph_or_style._p.get_or_add_pPr()
    else:
        ppr = paragraph_or_style._element.get_or_add_pPr()

    ind = ppr.get_or_add_ind()
    for name in (
        "left", "right", "firstLine",
        "leftChars", "rightChars", "firstLineChars",
        "start", "end", "startChars", "endChars",
    ):
        ind.set(qn(f"w:{name}"), "0")

    # firstLine and hanging are mutually exclusive; remove hanging rather than writing both.
    for name in ("hanging", "hangingChars"):
        ind.attrib.pop(qn(f"w:{name}"), None)


def iter_table_paragraphs(table):
    for row in table.rows:
        for cell in row.cells:
            yield from cell.paragraphs
            for nested in cell.tables:
                yield from iter_table_paragraphs(nested)


def iter_all_paragraphs(doc):
    yield from doc.paragraphs

    for table in doc.tables:
        yield from iter_table_paragraphs(table)

    for section in doc.sections:
        parts = (
            section.header,
            section.footer,
            section.first_page_header,
            section.first_page_footer,
            section.even_page_header,
            section.even_page_footer,
        )
        for part in parts:
            yield from part.paragraphs
            for table in part.tables:
                yield from iter_table_paragraphs(table)


def usable_width_twips(doc) -> int:
    widths = []
    for section in doc.sections:
        if (
            section.page_width is None
            or section.left_margin is None
            or section.right_margin is None
        ):
            continue
        usable = int(
            (section.page_width - section.left_margin - section.right_margin)
            / TWIP_EMU
        )
        if usable > 0:
            widths.append(usable)
    return min(widths) if widths else 8845


def normalize_table(table, table_style, max_width_twips: int) -> None:
    """Constrain width and normalize cell paragraph formatting."""
    table.autofit = False

    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is not None:
        tbl_w.set(qn("w:type"), "dxa")
        tbl_w.set(qn("w:w"), str(max_width_twips))

    # Do not add tblInd when it is absent; default is already zero.
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is not None:
        tbl_ind.set(qn("w:type"), "dxa")
        tbl_ind.set(qn("w:w"), "0")

    grid_cols = list(table._tbl.tblGrid.gridCol_lst)
    widths = []
    for col in grid_cols:
        try:
            widths.append(int(col.get(qn("w:w")) or "0"))
        except ValueError:
            widths.append(0)

    total = sum(widths)
    if widths and total > max_width_twips and total > 0:
        scaled = []
        used = 0
        for index, width in enumerate(widths):
            if index == len(widths) - 1:
                new_width = max_width_twips - used
            else:
                new_width = max(
                    1, round(width * max_width_twips / total)
                )
                used += new_width
            scaled.append(new_width)

        for col, width in zip(grid_cols, scaled):
            col.set(qn("w:w"), str(width))

    # Merged cells can be returned repeatedly by python-docx; process each tc only once.
    seen_cells = set()
    for row in table.rows:
        for cell in row.cells:
            key = id(cell._tc)
            if key in seen_cells:
                continue
            seen_cells.add(key)

            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

            # Let tblGrid control width so a long cell cannot expand the whole table.
            tcw = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            tcw.set(qn("w:type"), "auto")
            tcw.set(qn("w:w"), "0")

            for paragraph in cell.paragraphs:
                paragraph.style = table_style
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                set_indentation_zero(paragraph)
                set_pagination_off(paragraph.paragraph_format)

                for run in paragraph.runs:
                    if run.text:
                        run.font.size = Pt(TABLE_FONT_PT)

            for nested in cell.tables:
                normalize_table(nested, table_style, max_width_twips)


def verify_output(path: Path) -> None:
    """Fail before replacement if the generated package is structurally unreadable."""
    with zipfile.ZipFile(path, "r") as zf:
        bad_member = zf.testzip()
        if bad_member:
            raise RuntimeError(
                f"corrupt ZIP member after normalization: {bad_member}"
            )

        for name in zf.namelist():
            if name.endswith((".xml", ".rels")):
                try:
                    etree.fromstring(zf.read(name))
                except etree.XMLSyntaxError as exc:
                    raise RuntimeError(
                        f"invalid OOXML part after normalization: {name}: {exc}"
                    ) from exc

    # A second python-docx open catches package/relationship/content-type failures.
    Document(path)


def normalize(src: Path, dst: Path) -> None:
    doc = Document(src)

    try:
        body_style = doc.styles[BODY_STYLE]
    except KeyError as exc:
        raise RuntimeError(f"missing required style: {BODY_STYLE}") from exc

    body_style.paragraph_format.first_line_indent = Pt(BODY_FIRST_LINE_PT)

    try:
        table_style = doc.styles[TABLE_STYLE]
    except KeyError:
        table_style = doc.styles.add_style(
            TABLE_STYLE, WD_STYLE_TYPE.PARAGRAPH
        )
        table_style.base_style = body_style

    table_style.font.size = Pt(TABLE_FONT_PT)
    table_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_indentation_zero(table_style)
    set_pagination_off(table_style.paragraph_format)

    # Explicitly disable all four pagination options on every paragraph style.
    for style in doc.styles:
        if style.type == WD_STYLE_TYPE.PARAGRAPH:
            set_pagination_off(style.paragraph_format)

    # Main-body Normal prose becomes the canonical body style.
    # doc.paragraphs does not include table-cell paragraphs.
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        style_name = (
            paragraph.style.name
            if paragraph.style is not None
            else "Normal"
        )
        if (
            len(text) >= 12
            and paragraph.alignment
            not in (WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT)
            and style_name in {"Normal", BODY_STYLE}
        ):
            paragraph.style = body_style
            # Remove direct indentation so 正文（默认） supplies 640 twips.
            paragraph.paragraph_format.left_indent = None
            paragraph.paragraph_format.right_indent = None
            paragraph.paragraph_format.first_line_indent = None

    # Direct paragraph formatting must not re-enable pagination.
    for paragraph in iter_all_paragraphs(doc):
        set_pagination_off(paragraph.paragraph_format)

    max_width = usable_width_twips(doc)
    for table in doc.tables:
        normalize_table(table, table_style, max_width)

    dst.parent.mkdir(parents=True, exist_ok=True)
    doc.save(dst)
    verify_output(dst)


def main() -> int:
    parser = argparse.ArgumentParser(
        description=(
            "Safely normalize body indentation, pagination options, "
            "and table formatting in DOCX."
        )
    )
    parser.add_argument("document", type=Path)
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument("--in-place", action="store_true")
    group.add_argument("--out", type=Path)
    args = parser.parse_args()

    if args.in_place:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir) / args.document.name
            normalize(args.document, temp_path)
            # Only replace the original after every verification step has passed.
            os.replace(temp_path, args.document)
        target = args.document
    else:
        normalize(args.document, args.out)
        target = args.out

    print(f"Normalized and verified: {target}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
